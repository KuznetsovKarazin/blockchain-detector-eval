"""Build the manuscript with editable Word math; render_docx.py remains the QA gate."""
from pathlib import Path
import argparse,subprocess,re,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ap=argparse.ArgumentParser();ap.add_argument('source',type=Path);ap.add_argument('output',type=Path);a=ap.parse_args()
assert not re.search(r'PENDING|TODO|FIXME',a.source.read_text())
a.output.parent.mkdir(parents=True,exist_ok=True)
subprocess.run(['pandoc',str(a.source),'--from=markdown+tex_math_single_backslash+tex_math_dollars-implicit_figures','--to=docx','--resource-path='+str(a.source.parent),'--output='+str(a.output)],check=True)
doc=Document(a.output);styles={s.name:s for s in doc.styles};sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.3)
for st in doc.styles:
 if st.type in (1,2):
  st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
  for rf in st.element.findall('.//'+qn('w:rFonts')):
   for attr in list(rf.attrib):
    if attr.endswith('Theme'):del rf.attrib[attr]
  for color in st.element.findall('.//'+qn('w:color')):
   for attr in list(color.attrib):
    if 'theme' in attr:del color.attrib[attr]
for name in ['Normal','Body Text','First Paragraph']:
 st=styles[name];st.font.size=Pt(11);pf=st.paragraph_format
 pf.line_spacing=1.08;pf.space_after=Pt(5);pf.widow_control=True
for name,size in [('Title',18),('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
 st=styles[name];st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
for name in ['Verbatim Char','Source Code']:
 if name in styles:styles[name].font.name='Consolas';styles[name].font.size=Pt(9)
doc.paragraphs[0].style=styles['Title'];doc.paragraphs[0].paragraph_format.space_before=Pt(0)
in_refs=False
for p in doc.paragraphs:
 if p.style.name=='Heading 2':p.style=styles['Heading 1']
 elif p.style.name=='Heading 3':p.style=styles['Heading 2']
 if p.text.startswith('Table '):
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
  for r in p.runs:r.font.size=Pt(10)
 if p.text.startswith('Figure '):
  p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(8)
  for r in p.runs:r.font.size=Pt(10)
 if p._p.xpath('.//w:drawing'):
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(5)
 if p.text=='References':in_refs=True
 if in_refs:
  p.paragraph_format.keep_together=True
  if re.match(r'^\[\d+\]',p.text):
   p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(7)
   spacing=OxmlElement('w:contextualSpacing');spacing.set(qn('w:val'),'0')
   p._p.get_or_add_pPr().append(spacing)
for shape in doc.inline_shapes:
 w,h=shape.width,shape.height;shape.width=Inches(6.65);shape.height=int(h*shape.width/w)
widths=[[1.18,.75,2,2.97],[1.55,5.35],[1.28,1.1,1.25,1.25,2.02],[1.5,.84,1.17,.9,1.03,1.46],[1.5,1.42,1.42,1.42,1.14]]
for ti,t in enumerate(doc.tables):
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 ws=widths[ti];ws=[w*6.9/sum(ws) for w in ws]
 pr=t._tbl.tblPr
 for old in list(pr.findall(qn('w:tblBorders'))):pr.remove(old)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+side)
  for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
  borders.append(b)
 pr.append(borders);margins=OxmlElement('w:tblCellMar')
 for side,val in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
  m=OxmlElement('w:'+side);m.set(qn('w:w'),val);m.set(qn('w:type'),'dxa');margins.append(m)
 pr.append(margins)
 for c,w in zip(t.columns,ws):c.width=Inches(w)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(ws[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;prc=c._tc.get_or_add_tcPr()
   for old in list(prc.findall(qn('w:tcBorders'))):prc.remove(old)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F5');prc.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=False;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or ti in [0,1] else WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(9.5);r.font.bold=ri==0;r.font.color.rgb=RGBColor(0,0,0)
for node in doc.element.findall('.//'+qn('m:t')):
 if node.text==chr(92):node.text='∖'
 if node.text=='⌀':node.text='∅'
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();r.font.size=Pt(9);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
doc.core_properties.author='Oleksandr Kuznetsov';doc.core_properties.title=doc.paragraphs[0].text;doc.core_properties.comments=''
doc.save(a.output)
print(json.dumps({'output':str(a.output),'source_sha256':hashlib.sha256(a.source.read_bytes()).hexdigest(),'figures':len(doc.inline_shapes),'tables':len(doc.tables)}))
